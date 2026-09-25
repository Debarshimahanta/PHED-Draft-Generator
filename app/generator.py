from __future__ import annotations

import html
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .config import ASSETS, OUTPUTS, OFFICE_HEADER, COPY_CATEGORIES

A4_WIDTH = Cm(21.0)
A4_HEIGHT = Cm(29.7)


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = f"w:{edge}"
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in kwargs[edge]:
                element.set(qn(f"w:{key}"), str(kwargs[edge][key]))


def _font(run, *, bold=False, size=11, name="Times New Roman", italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _para(
    p,
    *,
    before=0,
    after=0,
    line=1.0,
    align=None,
    first_line_cm=None,
    left_cm=None,
    right_cm=None,
):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if first_line_cm is not None:
        pf.first_line_indent = Cm(first_line_cm)
    if left_cm is not None:
        pf.left_indent = Cm(left_cm)
    if right_cm is not None:
        pf.right_indent = Cm(right_cm)


def _copy_sentence(item: dict) -> str:
    designation = (item.get("designation") or "").strip()
    category = item.get("category") or "Custom"
    extra = (item.get("extra") or "").strip()
    ps_target = (item.get("ps_target") or "").strip()
    custom = (item.get("custom") or "").strip()

    if not designation:
        return ""

    if category == "PS Intimation":
        if ps_target:
            base = f"The {designation} for kind appraisal of the same to {ps_target}."
        else:
            base = f"The {designation} for kind intimation."
    elif category == "Custom":
        base = custom if custom else f"The {designation}"
    else:
        phrase = COPY_CATEGORIES.get(category, "") or ""
        base = f"The {designation} {phrase}".strip()

    if extra:
        if base and not base.endswith((".", "!", "?")):
            base += "."
        base += " " + extra
    return base.strip()


def _add_header(doc: Document):
    emblem = ASSETS / "assam_emblem.jpeg"
    if emblem.exists():
        p = doc.add_paragraph()
        _para(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run()
        r.add_picture(str(emblem), width=Cm(1.25))

    # Compact formal office header, matching the reference sample.
    for idx, line in enumerate(OFFICE_HEADER):
        p = doc.add_paragraph()
        _para(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(line)
        if idx == 0:
            _font(r, bold=True, size=10.5)
        elif idx == 1:
            _font(r, bold=True, size=10.5)
        elif idx == 2:
            _font(r, bold=True, size=10)
        else:
            _font(r, bold=True, size=9.5)

    # Reference format has no decorative rule under the letterhead.


def _add_signatory_text(doc: Document, *, signed_label=False):
    # Keep the designation block on the right side of the page, but center the
    # two designation lines within that block. This matches the office seal style.
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(10.8)
    table.columns[1].width = Cm(7.0)

    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_border(
            cell,
            top={"val": "nil"},
            left={"val": "nil"},
            bottom={"val": "nil"},
            right={"val": "nil"},
        )

    cell = table.cell(0, 1)
    p = cell.paragraphs[0]
    _para(p, before=7 if not signed_label else 5, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

    if signed_label:
        r = p.add_run("-SIGNED-")
        _font(r, bold=True, size=9.5)
        p = cell.add_paragraph()
        _para(p, before=2, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

    r = p.add_run("Chief Engineer (PHE) Water, Assam")
    _font(r, bold=True, size=10.5)

    p = cell.add_paragraph()
    _para(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Hengrabari, Guwahati - 36")
    _font(r, bold=True, size=10.5)


def build_docx(data: dict, output_path: Path | None = None) -> Path:
    OUTPUTS.mkdir(exist_ok=True)
    output_path = output_path or OUTPUTS / "generated_letter.docx"

    doc = Document()
    section = doc.sections[0]
    section.page_width = A4_WIDTH
    section.page_height = A4_HEIGHT
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.3)
    section.footer_distance = Cm(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    _add_header(doc)

    # Compact gap below header.
    p = doc.add_paragraph()
    _para(p, after=2)

    meta = doc.add_table(rows=1, cols=2)
    meta.autofit = False
    for cell in meta.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_border(
            cell,
            top={"val": "nil"},
            left={"val": "nil"},
            bottom={"val": "nil"},
            right={"val": "nil"},
        )

    p = meta.cell(0, 0).paragraphs[0]
    _para(p, after=0)
    r = p.add_run("No. ")
    _font(r, bold=True, size=10.5)
    r = p.add_run(data.get("file_no") or "")
    _font(r, bold=True, size=10.5)

    p = meta.cell(0, 1).paragraphs[0]
    _para(p, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    letter_date = (data.get("date") or "").strip()
    if letter_date:
        r = p.add_run(f"Dated: {letter_date}")
        _font(r, size=10)

    p = doc.add_paragraph()
    _para(p, before=4, after=0)
    r = p.add_run("To,")
    _font(r, bold=True, size=10.5)

    for line in (data.get("addressee") or "").splitlines():
        p = doc.add_paragraph()
        _para(p, after=0)
        r = p.add_run(line)
        _font(r, size=10.5)

    p = doc.add_paragraph()
    _para(p, before=7, after=3)
    r = p.add_run("Sub: ")
    _font(r, bold=True, size=10.5)
    r = p.add_run(data.get("subject") or "")
    _font(r, size=10.5)

    # Optional reference line if later added to data model.
    reference = (data.get("reference") or "").strip()
    if reference:
        p = doc.add_paragraph()
        _para(p, after=3)
        r = p.add_run("Ref: ")
        _font(r, bold=True, size=10.5)
        r = p.add_run(reference)
        _font(r, size=10.5)

    p = doc.add_paragraph()
    _para(p, before=4, after=4)
    r = p.add_run("Sir,")
    _font(r, size=10.5)

    body = data.get("body") or ""
    blocks = re.split(r"\n\s*\n", body)

    for block in blocks:
        if not block:
            continue
        p = doc.add_paragraph()
        _para(
            p,
            after=5,
            line=1.0,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line_cm=0.8,
        )
        r = p.add_run(block)
        _font(r, size=10.5)

    # First designation block after body, exactly as in the approved reference format.
    _add_signatory_text(doc, signed_label=False)

    copies = [x for x in data.get("copies", []) if (x.get("designation") or "").strip()]
    if copies:
        p = doc.add_paragraph()
        _para(p, before=8, after=3)
        r = p.add_run("Copy to:")
        _font(r, bold=True, size=10.5)

        for i, item in enumerate(copies, start=1):
            text = _copy_sentence(item)
            p = doc.add_paragraph()
            _para(
                p,
                after=1,
                line=1.0,
                align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                left_cm=0.65,
            )
            p.paragraph_format.first_line_indent = Cm(-0.42)
            r = p.add_run(f"{i}.  {text}")
            _font(r, size=9.7)

        # Final signed designation after Copy To.
        _add_signatory_text(doc, signed_label=True)

    doc.save(output_path)
    return output_path


def build_pdf(docx_path: Path) -> Path | None:
    out_dir = docx_path.parent
    commands = [
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
    ]

    for cmd in commands:
        try:
            subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
            pdf = out_dir / f"{docx_path.stem}.pdf"
            if pdf.exists():
                return pdf
        except Exception:
            continue
    return None


def preview_html(data: dict) -> str:
    def esc(value):
        return html.escape(value or "").replace("\n", "<br>")

    body = data.get("body") or ""
    body_html = "".join(
        f'<p>{html.escape(block).replace(chr(10), "<br>")}</p>'
        for block in re.split(r"\n\s*\n", body)
        if block
    )

    copies = []
    for i, item in enumerate(data.get("copies", []), start=1):
        if (item.get("designation") or "").strip():
            copies.append(
                f'<div class="copy"><span>{i}.</span><div>{html.escape(_copy_sentence(item))}</div></div>'
            )
    copy_html = "".join(copies)

    reference_html = ""
    if (data.get("reference") or "").strip():
        reference_html = f'<div class="reference"><b>Ref:</b> {esc(data.get("reference"))}</div>'

    return f'''<!doctype html>
<html>
<head>
<meta charset="utf-8">
<style>
*{{box-sizing:border-box}}
body{{font-family:"Times New Roman",serif;background:#eceff2;margin:0;padding:22px;color:#111}}
.page{{
    background:#fff;
    width:794px;
    min-height:1123px;
    margin:auto;
    padding:38px 58px 42px;
    box-shadow:0 3px 16px rgba(0,0,0,.18);
}}
.emblem{{text-align:center;line-height:1}}
.emblem img{{width:46px;height:auto}}
.header{{text-align:center;font-weight:700;line-height:1.08;font-size:12px}}
.header .small{{font-size:11px}}
.meta{{display:flex;justify-content:space-between;font-size:12px;margin-top:14px}}
.to{{font-size:12px;margin-top:14px;line-height:1.15}}
.subject,.reference{{font-size:12px;margin-top:13px;line-height:1.18}}
.reference{{margin-top:4px}}
.salutation{{font-size:12px;margin-top:12px}}
.body p{{
    font-size:12px;
    line-height:1.18;
    text-align:justify;
    text-indent:31px;
    margin:7px 0;
}}
.sign{{text-align:right;font-weight:700;font-size:11.5px;line-height:1.15;margin-top:17px}}
.copyhead{{font-size:12px;font-weight:700;margin-top:17px;margin-bottom:5px}}
.copy{{display:grid;grid-template-columns:20px 1fr;gap:4px;margin:2px 9px;font-size:11px;line-height:1.16;text-align:justify}}
.signed{{text-align:right;font-weight:700;font-size:10.5px;margin-top:20px}}
</style>
</head>
<body>
<div class="page">
    <div class="emblem"><img src="data:image/jpeg;base64,{{EMBLEM}}"></div>
    <div class="header">
        GOVERNMENT OF ASSAM<br>
        OFFICE OF THE CHIEF ENGINEER (PHE) WATER, ASSAM<br>
        HENGRABARI, GUWAHATI - 36<br>
        <span class="small">Email - asphe@rediffmail.com</span>
    </div>

    <div class="meta">
        <div><b>No. {esc(data.get("file_no"))}</b></div>
        <div>{'Dated: ' + esc(data.get("date")) if data.get("date") else ''}</div>
    </div>

    <div class="to"><b>To,</b><br>{esc(data.get("addressee"))}</div>
    <div class="subject"><b>Sub:</b> {esc(data.get("subject"))}</div>
    {reference_html}
    <div class="salutation">Sir,</div>
    <div class="body">{body_html}</div>

    <div class="signbox">Chief Engineer (PHE) Water, Assam<br>Hengrabari, Guwahati - 36</div>

    {'<div class="copyhead">Copy to:</div>' + copy_html if copy_html else ''}

    {('<div class="signbox signedbox"><div class="signed">-SIGNED-</div>Chief Engineer (PHE) Water, Assam<br>Hengrabari, Guwahati - 36</div>') if copy_html else ''}
</div>
</body>
</html>'''